from __future__ import annotations

from dataclasses import replace
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.core.card_models import CardTransaction
from app.core.enums import CardMatchStatus, EvidenceType, ReviewStatus
from app.core.ocr_models import OCRResult
from app.core.receipt_models import ReceiptRecord
from app.exporter.excel_exporter import ExcelExporter


REQUIRED_SHEETS = {
    "01_정리내역",
    "02_월별요약",
    "03_비용항목별요약",
    "04_검토필요",
    "05_카드대조",
    "06_원본파일목록",
    "07_처리로그",
    "08_제외내역",
}


def make_receipt(
    *,
    merchant_name: str = "테스트카페",
    transaction_date: date | None = date(2026, 5, 7),
    total_amount: int | None = 12000,
    supply_amount: int | None = 10909,
    vat_amount: int | None = 1091,
    expense_category: str | None = "복리후생비",
    recommended_expense_category: str | None = "복리후생비",
    review_status: str = ReviewStatus.APPROVED.value,
    card_match_status: str = CardMatchStatus.MATCHED.value,
    matched_card_transaction_id: str | None = None,
    card_match_score: float | None = 1.0,
    card_match_reason: str | None = "exact_date_amount_merchant",
    review_reasons: list[str] | None = None,
    validation_warnings: list[str] | None = None,
    parser_warnings: list[str] | None = None,
    duplicate_reason: str | None = None,
    export_included: bool = True,
    source_file_name: str = r"C:\private\receipt.jpg",
    user_note: str | None = None,
    gross_amount: int | None = None,
    discount_amount: int | None = None,
    final_amount: int | None = None,
    paid_amount: int | None = None,
    amount_source: str | None = None,
    amount_warnings: list[str] | None = None,
    estimated_supply_amount: int | None = None,
    estimated_vat_amount: int | None = None,
    vat_source: str | None = None,
    vat_note: str | None = None,
    vat_warnings: list[str] | None = None,
    vat_review_required: bool = False,
    sensitive_text_detected: bool = False,
    sensitive_text_categories: list[str] | None = None,
    sensitive_warnings: list[str] | None = None,
) -> ReceiptRecord:
    return ReceiptRecord.create(
        session_id="session-1",
        source_file_id=f"source-{merchant_name}",
        ocr_result_id=f"ocr-{merchant_name}",
        source_file_name=source_file_name,
        stored_relative_path="data/uploads/session-1/receipt_stored.jpg",
        evidence_type=EvidenceType.RECEIPT_IMAGE.value,
        transaction_date=transaction_date,
        merchant_name=merchant_name,
        business_number="123-45-67890",
        total_amount=total_amount,
        supply_amount=supply_amount,
        vat_amount=vat_amount,
        payment_method="신용카드",
        expense_category=expense_category,
        review_status=review_status,
        review_reason="; ".join(review_reasons or []),
        confidence=0.94,
        card_match_status=card_match_status,
        parser_warnings=parser_warnings or [],
        review_reasons=review_reasons or [],
        recommended_expense_category=recommended_expense_category,
        validation_warnings=validation_warnings or [],
        matched_card_transaction_id=matched_card_transaction_id,
        card_match_score=card_match_score,
        card_match_reason=card_match_reason,
        duplicate_reason=duplicate_reason,
        export_included=export_included,
        user_note=user_note,
        gross_amount=gross_amount,
        discount_amount=discount_amount,
        final_amount=final_amount,
        paid_amount=paid_amount,
        amount_source=amount_source,
        amount_warnings=amount_warnings or [],
        estimated_supply_amount=estimated_supply_amount,
        estimated_vat_amount=estimated_vat_amount,
        vat_source=vat_source,
        vat_note=vat_note,
        vat_warnings=vat_warnings or [],
        vat_review_required=vat_review_required,
        sensitive_text_detected=sensitive_text_detected,
        sensitive_text_categories=sensitive_text_categories or [],
        sensitive_warnings=sensitive_warnings or [],
    )


def make_card_transaction(*, transaction_id: str | None = None) -> CardTransaction:
    transaction = CardTransaction.create(
        session_id="session-1",
        source_file_id="card-source-1",
        source_file_name=r"C:\private\card.csv",
        stored_relative_path="data/uploads/session-1/card.csv",
        row_number=2,
        transaction_date=date(2026, 5, 7),
        merchant_name="테스트카페",
        total_amount=12000,
        card_name="테스트카드",
        approval_number="001234",
        currency="KRW",
        warning_messages=[],
    )
    if transaction_id is None:
        return transaction
    return replace(transaction, id=transaction_id)


def workbook_from_bytes(excel_bytes: bytes):
    return load_workbook(BytesIO(excel_bytes))


def headers(sheet) -> list[str]:
    return [cell.value for cell in sheet[1]]


def all_cell_values(workbook) -> list[object]:
    values = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            values.extend(cell.value for cell in row if cell.value is not None)
    return values


def test_excel_exporter_returns_xlsx_bytes_and_openpyxl_can_load():
    excel_bytes = ExcelExporter().export([make_receipt()])

    assert isinstance(excel_bytes, bytes)
    assert excel_bytes[:2] == b"PK"
    assert workbook_from_bytes(excel_bytes).sheetnames


def test_excel_exporter_saves_xlsx_file_when_output_path_is_given(tmp_path):
    output_path = tmp_path / "receipt_report.xlsx"

    saved_path = ExcelExporter().export([make_receipt()], output_path=output_path)

    assert saved_path == str(output_path)
    assert output_path.exists()
    assert load_workbook(output_path).sheetnames


def test_excel_exporter_creates_required_sheets():
    workbook = workbook_from_bytes(ExcelExporter().export([make_receipt()]))

    assert REQUIRED_SHEETS.issubset(set(workbook.sheetnames))


def test_summary_sheet_has_required_headers():
    workbook = workbook_from_bytes(ExcelExporter().export([make_receipt()]))

    assert headers(workbook["01_정리내역"]) == ExcelExporter.summary_headers


def test_summary_excludes_excluded_records_and_excluded_sheet_keeps_them():
    included = make_receipt(merchant_name="포함상호", total_amount=10000)
    excluded = make_receipt(
        merchant_name="제외상호",
        total_amount=20000,
        review_status=ReviewStatus.EXCLUDED.value,
        export_included=False,
        user_note="사용자가 확인 필요",
    )

    workbook = workbook_from_bytes(ExcelExporter().export([included, excluded]))

    summary = workbook["01_정리내역"]
    excluded_sheet = workbook["08_제외내역"]
    assert summary.max_row == 2
    assert summary["C2"].value == "포함상호"
    assert excluded_sheet.max_row == 2
    assert excluded_sheet["C2"].value == "제외상호"


def test_amount_and_date_cells_are_stored_as_values_with_formats():
    workbook = workbook_from_bytes(ExcelExporter().export([make_receipt()]))
    summary = workbook["01_정리내역"]

    assert summary["B2"].value.date() == date(2026, 5, 7)
    assert summary["B2"].number_format == "yyyy-mm-dd"
    assert isinstance(summary["E2"].value, int)
    assert summary["E2"].number_format == "#,##0"


def test_summary_sheet_includes_discount_amount_columns():
    receipt = make_receipt(
        total_amount=0,
        gross_amount=10800,
        discount_amount=10800,
        final_amount=0,
        paid_amount=0,
        amount_source="paid_amount",
        amount_warnings=["discount_detected", "zero_payment_after_discount"],
    )

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))
    summary = workbook["01_정리내역"]

    assert "할인전금액" in headers(summary)
    assert "금액검토사유" in headers(summary)
    assert summary["E2"].value == 0
    assert summary["F2"].value == 10800
    assert summary["G2"].value == 10800
    assert summary["H2"].value == 0
    assert summary["I2"].value == 0
    assert summary["J2"].value == "paid_amount"
    assert "zero_payment_after_discount" in summary["K2"].value
    assert summary["F2"].number_format == "#,##0"


def test_discount_zero_payment_pipeline_trace_excel_preserves_amount_state_without_raw_ocr_values():
    receipt = make_receipt(
        merchant_name="할인추적상호",
        source_file_name="sample_receipt.jpg",
        total_amount=0,
        gross_amount=10800,
        discount_amount=10800,
        final_amount=0,
        paid_amount=0,
        amount_source="paid_amount",
        amount_warnings=["discount_detected", "zero_payment_after_discount"],
        review_status=ReviewStatus.NEEDS_REVIEW.value,
        review_reasons=["할인 후 결제금액 0원, 사용자 확인 필요"],
    )
    ocr_result = OCRResult.completed(
        source_file_id="source-1",
        source_file_name="sample_receipt.jpg",
        stored_relative_path="data/uploads/session-1/sample_receipt.jpg",
        provider="mock",
        raw_text="RAW_TRACE_TEXT_SHOULD_NOT_APPEAR",
        raw_json={"marker": "RAW_TRACE_JSON_SHOULD_NOT_APPEAR"},
        confidence=0.9,
    )

    workbook = workbook_from_bytes(ExcelExporter().export_to_bytes([receipt], [ocr_result]))
    summary = workbook["01_정리내역"]
    review = workbook["04_검토필요"]
    column = {name: index + 1 for index, name in enumerate(headers(summary))}
    serialized = "\n".join(str(value) for value in all_cell_values(workbook))

    assert summary.cell(row=2, column=column["총금액"]).value == 0
    assert summary.cell(row=2, column=column["할인전금액"]).value == 10800
    assert summary.cell(row=2, column=column["할인금액"]).value == 10800
    assert summary.cell(row=2, column=column["최종결제금액"]).value == 0
    assert summary.cell(row=2, column=column["실제결제금액"]).value == 0
    assert summary.cell(row=2, column=column["금액판단기준"]).value == "paid_amount"
    assert "zero_payment_after_discount" in summary.cell(row=2, column=column["금액검토사유"]).value
    assert review.max_row == 2
    assert review["C2"].value == "할인추적상호"
    assert "RAW_TRACE_TEXT_SHOULD_NOT_APPEAR" not in serialized
    assert "RAW_TRACE_JSON_SHOULD_NOT_APPEAR" not in serialized


def test_summary_sheet_includes_estimated_vat_columns_with_reference_style():
    receipt = make_receipt(
        total_amount=5100,
        supply_amount=None,
        vat_amount=None,
        estimated_supply_amount=4636,
        estimated_vat_amount=464,
        vat_source="estimated_reference",
        vat_note="부가세 미표기, 참고용 역산값",
        vat_warnings=["vat_not_explicit", "vat_estimated_reference_only"],
        vat_review_required=True,
    )

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))
    summary = workbook["01_정리내역"]

    assert "추정공급가액" in headers(summary)
    assert "추정부가세" in headers(summary)
    assert summary["N2"].value == 4636
    assert summary["O2"].value == 464
    assert summary["P2"].value == "estimated_reference"
    assert "vat_estimated_reference_only" in summary["Q2"].value
    assert summary["N2"].number_format == "#,##0"
    assert summary["N1"].fill.fgColor.rgb == "00FFF2CC"
    assert summary["N2"].fill.fgColor.rgb == "00FFF2CC"


def test_summary_sheet_includes_sensitive_category_columns_without_values():
    receipt = make_receipt(
        sensitive_text_detected=True,
        sensitive_text_categories=["wifi_password", "card_number"],
        sensitive_warnings=["wifi_password_detected", "full_card_number_masked"],
        user_note="검수 메모",
    )

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))
    summary = workbook["01_정리내역"]
    serialized = "\n".join(str(value) for value in all_cell_values(workbook))

    assert "민감문구감지" in headers(summary)
    assert "민감문구유형" in headers(summary)
    assert summary["R2"].value == "Y"
    assert summary["S2"].value == "wifi_password; card_number"
    assert "sample-password-1234" not in serialized


def test_monthly_summary_uses_included_receipts_only():
    included = make_receipt(total_amount=10000, supply_amount=9000, vat_amount=1000)
    excluded = make_receipt(
        total_amount=20000,
        supply_amount=18000,
        vat_amount=2000,
        review_status=ReviewStatus.EXCLUDED.value,
        export_included=False,
    )

    workbook = workbook_from_bytes(ExcelExporter().export([included, excluded]))
    monthly = workbook["02_월별요약"]

    assert monthly["A2"].value == "2026-05"
    assert monthly["B2"].value == 1
    assert monthly["C2"].value == 10000
    assert monthly["D2"].value == 9000
    assert monthly["E2"].value == 1000


def test_category_summary_uses_expense_category_not_recommendation():
    receipt = make_receipt(
        expense_category="사용자입력",
        recommended_expense_category="시스템추천",
        total_amount=10000,
    )

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))
    category = workbook["03_비용항목별요약"]

    assert category["A2"].value == "사용자입력"
    assert category["C2"].value == 10000


def test_empty_expense_category_is_grouped_as_unconfirmed():
    receipt = make_receipt(expense_category=None, recommended_expense_category="시스템추천")

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))

    assert workbook["03_비용항목별요약"]["A2"].value == "미확정"


def test_recommended_category_does_not_fill_final_category_in_summary():
    receipt = make_receipt(expense_category=None, recommended_expense_category="시스템추천")

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))
    summary = workbook["01_정리내역"]

    assert summary["U2"].value == "시스템추천"
    assert summary["V2"].value is None


def test_review_sheet_includes_needs_review_and_warning_records():
    needs_review = make_receipt(
        merchant_name="검토상호",
        review_status=ReviewStatus.NEEDS_REVIEW.value,
        review_reasons=["금액 확인 필요"],
    )
    warning_only = make_receipt(
        merchant_name="경고상호",
        validation_warnings=["결제수단 확인 필요"],
        card_match_status=CardMatchStatus.MATCHED.value,
    )

    workbook = workbook_from_bytes(ExcelExporter().export([needs_review, warning_only]))
    review = workbook["04_검토필요"]

    assert review.max_row == 3
    merchants = {review["C2"].value, review["C3"].value}
    assert merchants == {"검토상호", "경고상호"}


def test_review_sheet_includes_zero_payment_after_discount_record():
    receipt = make_receipt(
        merchant_name="할인상호",
        total_amount=0,
        gross_amount=10800,
        discount_amount=10800,
        final_amount=0,
        paid_amount=0,
        review_status=ReviewStatus.NEEDS_REVIEW.value,
        review_reasons=["zero_payment_after_discount"],
        amount_warnings=["discount_detected", "zero_payment_after_discount"],
    )

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))
    review = workbook["04_검토필요"]

    assert review.max_row == 2
    assert review["C2"].value == "할인상호"
    assert "zero_payment_after_discount" in review["F2"].value


def test_review_sheet_includes_estimated_vat_reference_record():
    receipt = make_receipt(
        merchant_name="부가세미표기상호",
        total_amount=5100,
        supply_amount=None,
        vat_amount=None,
        estimated_supply_amount=4636,
        estimated_vat_amount=464,
        vat_source="estimated_reference",
        vat_warnings=["vat_not_explicit", "vat_estimated_reference_only"],
        vat_review_required=True,
    )

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))
    review = workbook["04_검토필요"]

    assert review.max_row == 2
    assert review["C2"].value == "부가세미표기상호"
    assert "vat_estimated_reference_only" in review["F2"].value


def test_review_sheet_includes_sensitive_text_detected_record_without_values():
    receipt = make_receipt(
        merchant_name="민감문구상호",
        sensitive_text_detected=True,
        sensitive_text_categories=["wifi_password"],
        sensitive_warnings=["wifi_password_detected"],
    )

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))
    review = workbook["04_검토필요"]
    serialized = "\n".join(str(value) for value in all_cell_values(workbook))

    assert review.max_row == 2
    assert review["C2"].value == "민감문구상호"
    assert "wifi_password_detected" in review["F2"].value
    assert "sample-password-1234" not in serialized


def test_card_match_status_and_reason_are_written():
    card = make_card_transaction(transaction_id="card-1")
    receipt = make_receipt(matched_card_transaction_id="card-1")

    workbook = workbook_from_bytes(ExcelExporter().export([receipt], [card]))
    sheet = workbook["05_카드대조"]

    assert sheet["E2"].value.date() == date(2026, 5, 7)
    assert sheet["F2"].value == "테스트카페"
    assert sheet["J2"].value == CardMatchStatus.MATCHED.value
    assert sheet["K2"].value == 1
    assert sheet["L2"].value == "exact_date_amount_merchant"


def test_unmatched_or_candidate_records_are_visible_in_review_or_card_sheet():
    receipt = make_receipt(
        card_match_status=CardMatchStatus.CANDIDATE.value,
        card_match_reason="date_plus_minus_1_amount_match",
        card_match_score=0.75,
    )

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))

    assert workbook["04_검토필요"]["H2"].value == CardMatchStatus.CANDIDATE.value
    assert workbook["05_카드대조"]["J2"].value == CardMatchStatus.CANDIDATE.value


def test_full_local_paths_are_not_written_to_any_sheet():
    receipt = make_receipt(source_file_name=r"C:\private\receipt.jpg")
    card = make_card_transaction()

    workbook = workbook_from_bytes(ExcelExporter().export([receipt], [card]))
    serialized = "\n".join(str(value) for value in all_cell_values(workbook))

    assert "C:" not in serialized
    assert "private" not in serialized
    assert "data/uploads" not in serialized


def test_ocr_raw_text_and_json_are_not_written_to_any_sheet():
    ocr_result = OCRResult.completed(
        source_file_id="source-1",
        source_file_name="receipt.jpg",
        stored_relative_path="data/uploads/session-1/receipt.jpg",
        provider="mock",
        raw_text="RAW_OCR_FULL_TEXT_SHOULD_NOT_APPEAR",
        raw_json={"field": "RAW_JSON_SHOULD_NOT_APPEAR"},
        confidence=0.9,
    )

    workbook = workbook_from_bytes(ExcelExporter().export_to_bytes([make_receipt()], [ocr_result]))
    serialized = "\n".join(str(value) for value in all_cell_values(workbook))

    assert "RAW_OCR_FULL_TEXT_SHOULD_NOT_APPEAR" not in serialized
    assert "RAW_JSON_SHOULD_NOT_APPEAR" not in serialized


def test_formula_like_strings_are_safely_stored_as_text():
    receipt = make_receipt(
        merchant_name="=SUM(1,2)",
        source_file_name="@receipt.jpg",
        user_note="+memo",
    )

    workbook = workbook_from_bytes(ExcelExporter().export([receipt]))
    summary = workbook["01_정리내역"]

    assert summary["C2"].value == "'=SUM(1,2)"
    assert summary["AB2"].value == "'@receipt.jpg"
    assert summary["AC2"].value == "'+memo"


def test_all_sheets_have_freeze_panes_and_auto_filter():
    workbook = workbook_from_bytes(ExcelExporter().export([make_receipt()]))

    for sheet in workbook.worksheets:
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None


def test_sample_output_schema_document_exists():
    assert Path("docs/sample_output_schema.md").exists()
