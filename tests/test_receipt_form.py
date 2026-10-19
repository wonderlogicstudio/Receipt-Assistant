from dataclasses import replace
from datetime import date

import pytest

from app.core.ocr_models import OCRResult
from app.parser.receipt_parser import ReceiptParser
from app.ui.receipt_form import apply_receipt_form, parse_review_amount
from app.ui.review_utils import merge_reviewed_records


def receipt():
    return ReceiptParser().parse(OCRResult.completed(
        source_file_id="synthetic", source_file_name="synthetic_receipt.png", stored_relative_path=None,
        provider="mock", raw_text="상호 테스트상점\n거래일자 2026-10-19\n합계 12000", raw_json={}, confidence=0.95,
    ))


def form_values():
    return dict(transaction_date=date(2026, 10, 19), merchant_name="검수상점", business_number="1234567890",
                discount_amount=0, supply_amount=None, vat_amount=None, total_amount=12000,
                payment_method="카드", review_status="검토필요", user_note="원본 확인 중")


def test_review_saves_business_number_discount_and_missing_taxes_without_changing_identity():
    original = replace(receipt(), estimated_supply_amount=10000, estimated_vat_amount=1000)
    updated = apply_receipt_form(original, form_values())
    assert updated.id == original.id and updated.source_file_id == original.source_file_id
    assert updated.source_file_name == original.source_file_name
    assert updated.merchant_name == "검수상점" and updated.business_number == "123-45-67890"
    assert updated.discount_amount == 0
    assert updated.supply_amount is None and updated.vat_amount is None
    assert updated.estimated_supply_amount is None and updated.estimated_vat_amount is None


def test_explicit_zero_and_unknown_amounts_remain_distinct():
    values = form_values()
    values.update(total_amount=0, supply_amount=0, vat_amount=0, discount_amount=None)
    result = apply_receipt_form(receipt(), values)
    assert result.total_amount == result.supply_amount == result.vat_amount == 0
    assert result.discount_amount is None


def test_duplicate_suspicion_does_not_automatically_exclude_or_delete():
    original = replace(receipt(), duplicate_group_id="synthetic-group", duplicate_candidate_ids=["other"])
    updated = apply_receipt_form(original, form_values())
    assert updated.export_included and updated.duplicate_candidate_ids == ["other"]
    values = form_values()
    values["review_status"] = "제외"
    assert not apply_receipt_form(original, values).export_included


def test_editing_selected_receipt_preserves_other_receipts_and_order():
    first, second = receipt(), receipt()
    updated = apply_receipt_form(second, form_values())
    merged = merge_reviewed_records([first, second], [updated])
    assert merged == [first, updated]


def test_manual_payment_edit_reaches_excel_final_column_without_rewriting_paid_amount():
    from io import BytesIO
    from openpyxl import load_workbook
    from app.exporter.excel_exporter import ExcelExporter
    original = replace(receipt(), total_amount=12000, final_amount=15000, paid_amount=12000)
    assert apply_receipt_form(original, form_values()).final_amount == 15000
    values = form_values()
    values["total_amount"] = 11000
    updated = apply_receipt_form(original, values)
    assert updated.total_amount == updated.final_amount == 11000
    assert updated.paid_amount == 12000
    sheet = load_workbook(BytesIO(ExcelExporter().export([updated])))[ExcelExporter.summary_sheet_name]
    columns = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, columns["최종결제금액"]).value == 11000
    assert sheet.cell(2, columns["실제결제금액"]).value == 12000


def review_app(tmp_path, monkeypatch, count=1):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setenv("APP_DB_PATH", str(tmp_path / "app.db"))
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    return AppTest.from_string('''
import streamlit as st
from app.core.ocr_models import OCRResult
from app.parser.receipt_parser import ReceiptParser
from app.ui.review_view import render_review_view
if "receipt_records" not in st.session_state:
    result = OCRResult.completed(source_file_id="synthetic", source_file_name="private_name.png", stored_relative_path=None,
        provider="mock", raw_text="Store\\n2026-10-19\\nTOTAL 12000", raw_json={"total_amount": 12000}, confidence=0.9)
    st.session_state["receipt_records"] = [ReceiptParser().parse(result) for _ in range(COUNT)]
render_review_view(st.session_state["receipt_records"])
'''.replace("COUNT", str(count)), default_timeout=20).run()


def button(app, label):
    return next(item for item in app.button if item.label == label)


def test_popup_has_native_close_control(tmp_path, monkeypatch):
    app = review_app(tmp_path, monkeypatch)
    button(app, "상세 검수 열기").click().run()
    assert not app.exception
    dialogs = app.get("dialog")
    assert len(dialogs) == 1
    assert dialogs[0].proto.dialog.dismissible


def text(app, label):
    return next(item for item in app.text_input if item.label == label)


def test_overview_first_then_popup_saves_review_and_invalidates_previous_export(tmp_path, monkeypatch):
    app = review_app(tmp_path, monkeypatch)
    assert not app.exception
    assert len(app.dataframe) == 1
    assert not app.text_input
    assert "private_name.png" not in str(app.selectbox(key="review_selected_receipt_id").options)
    button(app, "상세 검수 열기").click().run()
    assert not app.exception
    assert not any("후보 점수" in item.value or "하이라이트" in item.value for item in app.caption)
    app.session_state["export_file_bytes"] = b"previous export"
    text(app, "상호").set_value("수정한상호").run()
    text(app, "최종결제금액").set_value("12345").run()
    assert text(app, "최종결제금액").value == "12,345"
    assert app.session_state["receipt_records"][0].total_amount == 12000
    button(app, "저장").click().run()
    assert not app.exception
    assert app.session_state["receipt_records"][0].merchant_name == "수정한상호"
    assert app.session_state["receipt_records"][0].total_amount == 12345
    assert app.session_state["export_file_bytes"] is None
    assert app.dataframe[0].value.iloc[0]["최종결제금액"] == 12345


def test_popup_close_and_navigation_keep_unsaved_drafts_out_of_records(tmp_path, monkeypatch):
    app = review_app(tmp_path, monkeypatch, count=3)
    originals = list(app.session_state["receipt_records"])
    button(app, "상세 검수 열기").click().run()
    text(app, "상호").set_value("첫 번째 초안").run()
    button(app, "다음").click().run()
    assert text(app, "상호").value != "첫 번째 초안"
    text(app, "상호").set_value("두 번째 초안").run()
    button(app, "이전").click().run()
    assert text(app, "상호").value == "첫 번째 초안"
    assert button(app, "이전").disabled
    button(app, "닫기").click().run()
    assert not app.session_state["receipt_review_dialog_open"]
    assert app.session_state["receipt_records"] == originals
    drafts = app.session_state["receipt_review_drafts"]
    assert drafts[originals[0].id]["merchant_name"] == "첫 번째 초안"
    assert drafts[originals[1].id]["merchant_name"] == "두 번째 초안"
    assert not app.exception
    # Recreate client widgets after closing; AppTest retains a closed dialog tree.
    reopened = review_app(tmp_path, monkeypatch, count=3)
    reopened.session_state["receipt_records"] = originals
    reopened.session_state["receipt_review_drafts"] = drafts
    reopened.run()
    button(reopened, "상세 검수 열기").click().run()
    assert text(reopened, "상호").value == "첫 번째 초안"
    button(reopened, "다음").click().run()
    assert text(reopened, "상호").value == "두 번째 초안"


def test_invalid_amount_cannot_silently_clear_a_saved_value(tmp_path, monkeypatch):
    app = review_app(tmp_path, monkeypatch)
    button(app, "상세 검수 열기").click().run()
    text(app, "최종결제금액").set_value("12,34oops").run()
    button(app, "저장 후 다음").click().run()
    assert app.session_state["receipt_records"][0].total_amount == 12000
    assert app.session_state["receipt_review_dialog_open"]
    assert app.error and not app.exception


def test_save_next_updates_only_selected_record_and_closes_at_end(tmp_path, monkeypatch):
    app = review_app(tmp_path, monkeypatch, count=2)
    ids = [record.id for record in app.session_state["receipt_records"]]
    button(app, "상세 검수 열기").click().run()
    text(app, "상호").set_value("저장한 첫 영수증")
    button(app, "저장 후 다음").click().run()
    assert app.session_state["receipt_review_dialog_id"] == ids[1]
    assert app.session_state["receipt_records"][0].merchant_name == "저장한 첫 영수증"
    assert app.session_state["receipt_records"][1].merchant_name != "저장한 첫 영수증"
    assert button(app, "다음").disabled
    button(app, "저장 후 다음").click().run()
    assert not app.session_state["receipt_review_dialog_open"]
    assert not app.exception


def test_saving_with_review_filter_keeps_popup_navigation_valid(tmp_path, monkeypatch):
    app = review_app(tmp_path, monkeypatch, count=2)
    app.selectbox(key="receipt_review_status_filter").select("검토필요").run()
    button(app, "상세 검수 열기").click().run()
    next(item for item in app.selectbox if item.label == "검수상태").select("확인완료")
    button(app, "저장 후 다음").click().run()
    assert not app.exception
    assert len(app.dataframe[0].value) == 1
    button(app, "닫기").click().run()
    assert app.selectbox(key="receipt_review_status_filter").value == "검토필요"


def test_filter_changes_keep_the_selected_receipt_when_it_is_still_visible(tmp_path, monkeypatch):
    app = review_app(tmp_path, monkeypatch, count=3)
    records = list(app.session_state["receipt_records"])
    records[1] = replace(records[1], review_status="approved")
    app.session_state["receipt_records"] = records
    app.run()
    app.selectbox(key="review_selected_receipt_id").select(records[2].id).run()
    app.selectbox(key="receipt_review_status_filter").select("검토필요").run()
    assert app.selectbox(key="review_selected_receipt_id").value == records[2].id
    app.selectbox(key="receipt_review_status_filter").select("전체").run()
    assert app.selectbox(key="review_selected_receipt_id").value == records[2].id


def test_closing_popup_clears_grid_selection_for_same_row_reopening(tmp_path, monkeypatch):
    app = review_app(tmp_path, monkeypatch, count=2)
    original_grid_id = app.dataframe[0].proto.id
    assert app.session_state[app.dataframe[0].key]["selection"]["cells"] == []
    button(app, "상세 검수 열기").click().run()
    selected = app.session_state["receipt_review_dialog_id"]
    button(app, "닫기").click().run()
    assert app.dataframe[0].proto.id != original_grid_id
    assert app.session_state[app.dataframe[0].key]["selection"]["cells"] == []
    assert app.selectbox(key="review_selected_receipt_id").value == selected
    button(app, "상세 검수 열기").click().run()
    assert app.session_state["receipt_review_dialog_id"] == selected
    assert not app.exception


@pytest.mark.parametrize("value, expected", [("1,234", 1234), ("0", 0), ("", None), (None, None), ("—", None), ("-12,000", -12000)])
def test_review_money_preserves_commas_zero_missing_and_refunds(value, expected):
    assert parse_review_amount(value) == expected


@pytest.mark.parametrize("value", ["1,23", "abc", "12.34", True, "1e4", "12,000원", "9007199254740992"])
def test_review_money_rejects_ambiguous_or_invalid_input(value):
    with pytest.raises(ValueError):
        parse_review_amount(value)


def test_card_summary_distinguishes_missing_upload_from_zero_matches():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_string('''
import streamlit as st
from app.storage.file_store import StoredFileMetadata
from app.ui.streamlit_app import render_processing_summary
card = st.checkbox("synthetic card upload")
metadata = StoredFileMetadata("card.csv" if card else "receipt.png", "synthetic", "synthetic", ".csv" if card else ".png", 1, "synthetic")
render_processing_summary([metadata], [], [])
''', default_timeout=20).run()
    metrics = {item.label: item.value for item in app.metric}
    assert metrics["카드 사용내역 건수"] == "미업로드"
    assert metrics["카드매칭"] == metrics["카드 후보/미매칭"] == "미실행"
    app.checkbox[0].check().run()
    assert not app.exception
    metrics = {item.label: item.value for item in app.metric}
    assert metrics["카드 사용내역 건수"] == metrics["카드매칭"] == "0"
    assert metrics["카드 후보/미매칭"] == "0/0"


def test_gross_and_discount_do_not_generate_a_missing_final_amount():
    record = ReceiptParser().parse(OCRResult.completed(
        source_file_id="synthetic", source_file_name="synthetic.png", stored_relative_path=None,
        provider="mock", raw_text="상호 테스트상점\n총매출액 12000\n할인금액 2000", raw_json={}, confidence=0.9,
    ))
    assert record.gross_amount == 12000 and record.discount_amount == 2000
    assert record.final_amount is None and record.total_amount is None
    assert record.supply_amount is None and record.vat_amount is None
