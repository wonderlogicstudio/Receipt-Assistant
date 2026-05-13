from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.core.config import AppConfig
from app.core.enums import CardMatchStatus, ReviewStatus
from app.core.exceptions import InvalidSessionIdError
from app.core.ocr_models import OCRResult
from app.db.repository import JobHistoryRepository
from app.history.job_service import JobHistoryService
from app.storage.cleanup import cleanup_session_files
from app.storage.file_store import FileStore, StoredFileMetadata
from app.ui.history_utils import jobs_to_history_dataframe
from app.ui.review_utils import (
    COL_EXPENSE_CATEGORY,
    COL_MERCHANT,
    COL_REVIEW_STATUS,
    COL_TOTAL,
    STATUS_LABEL_EXCLUDED,
    receipt_records_to_review_dataframe,
    review_dataframe_to_receipt_records,
)
from app.ui.streamlit_app import (
    build_export_filename,
    save_export_history_once,
    prepare_export_bytes,
    process_card_statement_files,
    process_receipt_files_with_mock_ocr,
    process_uploaded_file_metadata,
    route_uploaded_files_by_extension,
)


class UploadFileStub:
    def __init__(self, name: str, data: bytes) -> None:
        self.name = name
        self._data = data

    def getvalue(self) -> bytes:
        return self._data


class FailingOCRProvider:
    provider_name = "azure"

    def extract(self, source):
        return OCRResult.failed(
            source_file_id="source-id",
            source_file_name="receipt.jpg",
            stored_relative_path=None,
            provider=self.provider_name,
            error_message="hidden internal detail",
            diagnostic_code="proxy_error",
        )


class ColumnarDiscountOCRProvider:
    provider_name = "azure"

    def extract(self, source):
        raw_text = "\n".join(
            [
                "상호명: 테스트상점",
                "사업자번호: 123-45-67890",
                "거래일자: 2026.05.10",
                "제품명",
                "레디팩(아몬드봉봉) 10,800",
                "KT 멤버십 VVIP 초이스 -10,800",
                "총매출액",
                "할인금액",
                "합계금액",
                "받은금액",
                "10,800",
                "10,800",
                "0",
                "0",
            ]
        )
        return OCRResult.completed(
            source_file_id="source-id",
            source_file_name="sample_receipt.jpg",
            stored_relative_path=None,
            provider=self.provider_name,
            raw_text=raw_text,
            raw_json={},
            confidence=0.94,
        )


def make_file_store(tmp_path: Path) -> FileStore:
    return FileStore(
        AppConfig(
            upload_root=tmp_path / "data" / "uploads",
            upload_relative_root=Path("data/uploads"),
        )
    )


def save_upload(file_store: FileStore, session_id: str, name: str, data: bytes = b"sample") -> StoredFileMetadata:
    return file_store.save_uploaded_file(UploadFileStub(name, data), session_id=session_id)


def card_csv_bytes(*, merchant: str = "테스트카페", amount: int = 12000, transaction_date: str = "2026-05-07") -> bytes:
    content = "\n".join(
        [
            "date,merchant,amount,card,approval_no",
            f"{transaction_date},{merchant},{amount},테스트카드,001234",
        ]
    )
    return content.encode("utf-8")


def card_xlsx_bytes(*, merchant: str = "테스트카페", amount: int = 12000) -> bytes:
    buffer = BytesIO()
    dataframe = pd.DataFrame(
        [
            {
                "date": "2026-05-07",
                "merchant": merchant,
                "amount": amount,
                "card": "테스트카드",
                "approval_no": "001234",
            }
        ]
    )
    dataframe.to_excel(buffer, index=False)
    return buffer.getvalue()


def workbook_from_export(receipts, card_transactions=None):
    return load_workbook(BytesIO(prepare_export_bytes(receipts, card_transactions or [])))


def all_cell_values(workbook) -> list[object]:
    values = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            values.extend(cell.value for cell in row if cell.value is not None)
    return values


def test_route_uploaded_files_by_extension_splits_receipts_cards_and_unsupported():
    metadata = [
        StoredFileMetadata("a.jpg", "a.jpg", "data/uploads/s/a.jpg", ".jpg", 1, "s"),
        StoredFileMetadata("b.jpeg", "b.jpeg", "data/uploads/s/b.jpeg", ".jpeg", 1, "s"),
        StoredFileMetadata("c.png", "c.png", "data/uploads/s/c.png", ".png", 1, "s"),
        StoredFileMetadata("d.pdf", "d.pdf", "data/uploads/s/d.pdf", ".pdf", 1, "s"),
        StoredFileMetadata("e.xlsx", "e.xlsx", "data/uploads/s/e.xlsx", ".xlsx", 1, "s"),
        StoredFileMetadata("f.csv", "f.csv", "data/uploads/s/f.csv", ".csv", 1, "s"),
        StoredFileMetadata("g.txt", "g.txt", "data/uploads/s/g.txt", ".txt", 1, "s"),
    ]

    routed = route_uploaded_files_by_extension(metadata)

    assert [item.file_extension for item in routed.receipt_files] == [".jpg", ".jpeg", ".png", ".pdf"]
    assert [item.file_extension for item in routed.card_statement_files] == [".xlsx", ".csv"]
    assert [item.file_extension for item in routed.unsupported_files] == [".txt"]


def test_process_receipt_files_with_mock_ocr_creates_validated_receipts(tmp_path):
    file_store = make_file_store(tmp_path)
    receipt_metadata = save_upload(file_store, "session1", "receipt.jpg")

    result = process_receipt_files_with_mock_ocr([receipt_metadata], file_store=file_store)

    assert len(result.receipt_records) == 1
    receipt = result.receipt_records[0]
    assert receipt.review_status in {ReviewStatus.PENDING.value, ReviewStatus.NEEDS_REVIEW.value}
    assert receipt.recommended_expense_category is not None
    assert receipt.source_file_name == "receipt.jpg"
    assert len(result.ocr_results) == 1


def test_process_receipt_failure_reports_safe_ocr_diagnostic(tmp_path):
    file_store = make_file_store(tmp_path)
    receipt_metadata = save_upload(file_store, "session1", "receipt.jpg")

    result = process_receipt_files_with_mock_ocr(
        [receipt_metadata],
        file_store=file_store,
        ocr_provider=FailingOCRProvider(),
    )

    assert result.receipt_records == []
    assert len(result.errors) == 1
    assert "status=failed, provider=azure, reason=proxy_error" in result.errors[0]
    assert "hidden internal detail" not in result.errors[0]
    assert str(tmp_path) not in result.errors[0]


def test_process_receipt_flow_keeps_columnar_discount_zero_payment_amounts(tmp_path):
    file_store = make_file_store(tmp_path)
    receipt_metadata = save_upload(file_store, "session1", "sample_receipt.jpg")

    result = process_receipt_files_with_mock_ocr(
        [receipt_metadata],
        file_store=file_store,
        ocr_provider=ColumnarDiscountOCRProvider(),
    )

    assert result.errors == []
    assert len(result.receipt_records) == 1
    receipt = result.receipt_records[0]
    assert receipt.gross_amount == 10800
    assert receipt.discount_amount == 10800
    assert receipt.final_amount == 0
    assert receipt.paid_amount == 0
    assert receipt.total_amount == 0
    assert receipt.amount_source == "paid_amount"
    assert "discount_detected" in receipt.amount_warnings
    assert "zero_payment_after_discount" in receipt.amount_warnings


def test_process_card_statement_files_reads_csv_and_xlsx(tmp_path):
    file_store = make_file_store(tmp_path)
    csv_metadata = save_upload(file_store, "session1", "card.csv", card_csv_bytes())
    xlsx_metadata = save_upload(file_store, "session1", "card.xlsx", card_xlsx_bytes())

    result = process_card_statement_files([csv_metadata, xlsx_metadata], file_store=file_store)

    assert len(result.card_transactions) == 2
    assert {transaction.source_file_name for transaction in result.card_transactions} == {"card.csv", "card.xlsx"}
    assert result.card_transactions[0].approval_number == "001234"


def test_full_processing_applies_duplicate_checker_and_card_matcher(tmp_path):
    file_store = make_file_store(tmp_path)
    first_receipt = save_upload(file_store, "session1", "receipt-a.jpg")
    second_receipt = save_upload(file_store, "session1", "receipt-b.jpg")
    receipt_result = process_receipt_files_with_mock_ocr([first_receipt], file_store=file_store)
    receipt = receipt_result.receipt_records[0]
    card_metadata = save_upload(
        file_store,
        "session1",
        "card.csv",
        card_csv_bytes(
            merchant=receipt.merchant_name or "테스트카페",
            amount=receipt.total_amount or 0,
            transaction_date=receipt.transaction_date.isoformat() if receipt.transaction_date else "2026-05-06",
        ),
    )

    result = process_uploaded_file_metadata([first_receipt, second_receipt, card_metadata], file_store=file_store)

    assert len(result.receipt_records) == 2
    assert all(record.duplicate_reason for record in result.receipt_records)
    assert all(record.card_match_status == CardMatchStatus.MATCHED.value for record in result.receipt_records)
    assert len(result.card_transactions) == 1


def test_export_flow_works_without_card_transactions(tmp_path):
    file_store = make_file_store(tmp_path)
    receipt_metadata = save_upload(file_store, "session1", "receipt.jpg")
    result = process_uploaded_file_metadata([receipt_metadata], file_store=file_store)

    workbook = workbook_from_export(result.receipt_records, [])

    assert "01_정리내역" in workbook.sheetnames
    assert workbook["05_카드대조"].max_row == 2


def test_reviewed_dataframe_changes_are_reflected_in_excel(tmp_path):
    file_store = make_file_store(tmp_path)
    receipt_metadata = save_upload(file_store, "session1", "receipt.jpg")
    result = process_uploaded_file_metadata([receipt_metadata], file_store=file_store)
    dataframe = receipt_records_to_review_dataframe(result.receipt_records)
    dataframe.loc[0, COL_MERCHANT] = "수정상호"
    dataframe.loc[0, COL_TOTAL] = 45000
    dataframe.loc[0, COL_EXPENSE_CATEGORY] = "사용자입력"

    reviewed = review_dataframe_to_receipt_records(dataframe, result.receipt_records)
    workbook = workbook_from_export(reviewed, result.card_transactions)

    summary = workbook["01_정리내역"]
    assert summary["C2"].value == "수정상호"
    assert summary["E2"].value == 45000
    assert summary["V2"].value == "사용자입력"


def test_excluded_reviewed_record_is_not_in_summary_but_is_kept_in_excluded_sheet(tmp_path):
    file_store = make_file_store(tmp_path)
    receipt_metadata = save_upload(file_store, "session1", "receipt.jpg")
    result = process_uploaded_file_metadata([receipt_metadata], file_store=file_store)
    dataframe = receipt_records_to_review_dataframe(result.receipt_records)
    dataframe.loc[0, COL_REVIEW_STATUS] = STATUS_LABEL_EXCLUDED

    reviewed = review_dataframe_to_receipt_records(dataframe, result.receipt_records)
    workbook = workbook_from_export(reviewed, result.card_transactions)

    assert workbook["01_정리내역"].max_row == 1
    assert workbook["08_제외내역"].max_row == 2


def test_build_export_filename_uses_expected_safe_format():
    filename = build_export_filename(datetime(2026, 5, 8, 14, 30, 0))

    assert filename == "receipt_summary_20260508_143000.xlsx"
    assert " " not in filename
    assert all(char.isascii() for char in filename)


def test_cleanup_session_files_dry_run_returns_candidates_without_deleting(tmp_path):
    config = AppConfig(upload_root=tmp_path / "data" / "uploads", upload_relative_root=Path("data/uploads"))
    session_dir = config.upload_root / "session1"
    session_dir.mkdir(parents=True)
    target = session_dir / "receipt.jpg"
    target.write_bytes(b"sample")

    summary = cleanup_session_files("session1", config=config, dry_run=True)

    assert summary["scanned_count"] == 1
    assert summary["eligible_count"] == 1
    assert summary["deleted_count"] == 0
    assert "receipt.jpg" not in str(summary)
    assert target.exists()


def test_cleanup_session_files_rejects_path_traversal(tmp_path):
    config = AppConfig(upload_root=tmp_path / "data" / "uploads", upload_relative_root=Path("data/uploads"))

    with pytest.raises(InvalidSessionIdError):
        cleanup_session_files("../outside", config=config, dry_run=True)


def test_export_does_not_include_local_paths_or_raw_ocr_text(tmp_path):
    file_store = make_file_store(tmp_path)
    receipt_metadata = save_upload(file_store, "session1", "receipt.jpg")
    result = process_uploaded_file_metadata([receipt_metadata], file_store=file_store)

    workbook = workbook_from_export(result.receipt_records, result.card_transactions)
    serialized = "\n".join(str(value) for value in all_cell_values(workbook))

    assert str(tmp_path) not in serialized
    assert "data/uploads" not in serialized
    assert "raw_text" not in serialized
    assert "raw_json" not in serialized


def test_excel_generation_persists_completed_job_history_once(tmp_path):
    config = AppConfig(
        upload_root=tmp_path / "data" / "uploads",
        upload_relative_root=Path("data/uploads"),
        output_root=tmp_path / "data" / "outputs",
        app_db_path=tmp_path / "history.db",
    )
    file_store = FileStore(config)
    service = JobHistoryService(repository=JobHistoryRepository(config=config), config=config)
    receipt_metadata = save_upload(file_store, "session1", "receipt.jpg")
    card_metadata = save_upload(file_store, "session1", "card.csv", card_csv_bytes())
    result = process_uploaded_file_metadata([receipt_metadata, card_metadata], file_store=file_store)
    export_bytes = prepare_export_bytes(result.receipt_records, result.card_transactions)
    session_state: dict[str, object] = {}

    job = save_export_history_once(
        session_state=session_state,
        session_id="session1",
        uploaded_files=result.uploaded_file_metadata,
        receipt_records=result.receipt_records,
        card_transactions=result.card_transactions,
        export_file_name="receipt_summary_20260510_120000.xlsx",
        export_file_bytes=export_bytes,
        service=service,
    )
    duplicate = save_export_history_once(
        session_state=session_state,
        session_id="session1",
        uploaded_files=result.uploaded_file_metadata,
        receipt_records=result.receipt_records,
        card_transactions=result.card_transactions,
        export_file_name="receipt_summary_20260510_120000.xlsx",
        export_file_bytes=export_bytes,
        service=service,
    )

    jobs = service.list_processing_jobs()
    detail = service.get_processing_job_detail(str(job["id"]))
    export_file = detail["export_files"][0]
    serialized_job = "\n".join(str(value) for value in job.values() if value is not None)

    assert duplicate["already_saved"] is True
    assert len(jobs) == 1
    assert jobs[0]["receipt_count"] == len(result.receipt_records)
    assert jobs[0]["card_transaction_count"] == len(result.card_transactions)
    assert jobs[0]["needs_review_count"] == sum(
        1 for receipt in result.receipt_records if receipt.review_status == ReviewStatus.NEEDS_REVIEW.value
    )
    assert jobs[0]["total_amount"] == sum(receipt.total_amount or 0 for receipt in result.receipt_records)
    assert export_file["export_file_name"] == "receipt_summary_20260510_120000.xlsx"
    assert export_file["export_relative_path"] == f"data/outputs/{job['id']}/receipt_summary_20260510_120000.xlsx"
    assert (config.output_root / str(job["id"]) / "receipt_summary_20260510_120000.xlsx").exists()
    assert str(tmp_path) not in serialized_job
    assert "raw_text" not in serialized_job
    assert "raw_json" not in serialized_job
    assert not jobs_to_history_dataframe(jobs).empty
