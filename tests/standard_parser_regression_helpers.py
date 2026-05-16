from __future__ import annotations

import json
import re
from dataclasses import asdict, replace
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from app.core.card_models import CardTransaction
from app.core.receipt_models import ReceiptRecord
from app.core.standard_ocr import StandardOCRResult
from app.ocr.adapters.azure_standard_adapter import AzureStandardOCRAdapter
from app.parser.receipt_parser import ReceiptParser
from app.parser.standard_ocr_context import build_parser_context_from_standard_ocr


FIXTURE_ROOT = Path(__file__).parent / "fixtures"
STANDARD_FIXTURE_DIR = FIXTURE_ROOT / "standard_ocr"
AZURE_FIXTURE_DIR = FIXTURE_ROOT / "azure"

STANDARD_FIXTURES = [
    "mock_receipt_standard.json",
    "azure_like_receipt_standard.json",
    "edge_case_receipt_standard.json",
]

AZURE_ADAPTER_FIXTURES = [
    "azure_receipt_sanitized_response.json",
    "azure_receipt_edge_sanitized_response.json",
    "azure_receipt_failed_sanitized_response.json",
]

FORBIDDEN_TEXT_PATTERNS = [
    r"raw_json",
    r"raw_text",
    r"provider_raw_response",
    r"full_ocr_text",
    r"endpoint",
    r"operation_location",
    r"api_key",
    r"secret",
    r"token",
    r"password",
    r"[A-Za-z]:\\",
    r"/Users/",
    r"/home/",
    r"/var/",
    r"/tmp/",
    r"\b\d{13,19}\b",
    r"\b\d{6}-\d{7}\b",
]

RAW_METADATA_MARKERS = [
    "provider_fields",
    "provider_family",
    "adapter_version",
    "normalized_block_count",
    "normalized_candidate_count",
    "analyze_result",
    "read_result",
]

AMOUNT_FIELDS = [
    "total_amount",
    "gross_amount",
    "discount_amount",
    "final_amount",
    "paid_amount",
    "supply_amount",
    "vat_amount",
]


def load_standard_result(name: str) -> StandardOCRResult:
    return StandardOCRResult.model_validate_json(
        (STANDARD_FIXTURE_DIR / name).read_text(encoding="utf-8")
    )


def convert_azure_fixture(name: str) -> StandardOCRResult:
    response = json.loads((AZURE_FIXTURE_DIR / name).read_text(encoding="utf-8"))
    return AzureStandardOCRAdapter().convert(
        response,
        source_file_id=f"azure-{name}",
        source_file_name="sample_azure_receipt.pdf",
    )


def parse_standard_result(result: StandardOCRResult) -> ReceiptRecord:
    context = build_parser_context_from_standard_ocr(result)
    return ReceiptParser().parse_standard_context(
        context,
        source_file_id=result.source_file_id,
        source_file_name=result.source_file_name,
        ocr_provider=result.provider,
        confidence=result.confidence,
    )


def parse_standard_fixture(name: str) -> ReceiptRecord:
    return parse_standard_result(load_standard_result(name))


def parse_azure_adapter_fixture(name: str) -> ReceiptRecord:
    return parse_standard_result(convert_azure_fixture(name))


def standard_parser_records() -> list[ReceiptRecord]:
    records = [parse_standard_fixture(name) for name in STANDARD_FIXTURES]
    records.extend(parse_azure_adapter_fixture(name) for name in AZURE_ADAPTER_FIXTURES)
    return records


def paddleocr_candidate_record(base: ReceiptRecord | None = None) -> ReceiptRecord:
    source = base or parse_standard_fixture("mock_receipt_standard.json")
    return replace(
        source,
        ocr_provider="paddleocr",
        source_file_id="paddleocr-candidate-source",
        source_file_name="sample_paddleocr_candidate.png",
    )


def matching_card_transaction(receipt: ReceiptRecord) -> CardTransaction:
    return CardTransaction.create(
        session_id=receipt.session_id,
        source_file_id="card-source-1",
        source_file_name="sample_card_statement.csv",
        stored_relative_path=None,
        row_number=2,
        transaction_date=receipt.transaction_date,
        merchant_name=receipt.merchant_name,
        total_amount=receipt.total_amount,
        card_name="sample-card",
        approval_number=receipt.approval_number or "APPROVAL-CARD-1",
        warning_messages=[],
    )


def workbook_from_bytes(excel_bytes: bytes):
    return load_workbook(BytesIO(excel_bytes))


def workbook_cell_values(workbook) -> list[object]:
    values: list[object] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            values.extend(cell.value for cell in row if cell.value is not None)
    return values


def serialize_for_scan(value: object) -> str:
    if isinstance(value, ReceiptRecord):
        value = value.to_dict()
    if isinstance(value, CardTransaction):
        value = asdict(value)
    return json.dumps(value, ensure_ascii=False, default=str)


def assert_no_forbidden_text(value: object) -> None:
    dumped = serialize_for_scan(value)
    for pattern in FORBIDDEN_TEXT_PATTERNS:
        assert not re.search(pattern, dumped, flags=re.IGNORECASE), pattern


def assert_no_raw_metadata(value: object) -> None:
    dumped = serialize_for_scan(value)
    for marker in RAW_METADATA_MARKERS:
        assert marker not in dumped


def assert_no_forbidden_workbook_values(values: Iterable[object]) -> None:
    dumped = "\n".join(str(value) for value in values if value is not None)
    for pattern in FORBIDDEN_TEXT_PATTERNS:
        assert not re.search(pattern, dumped, flags=re.IGNORECASE), pattern
    for marker in RAW_METADATA_MARKERS:
        assert marker not in dumped


def assert_non_amount_values_not_in_amount_fields(record: ReceiptRecord) -> None:
    amount_values = {getattr(record, field_name) for field_name in AMOUNT_FIELDS}
    if record.business_number is not None:
        assert record.business_number not in amount_values
    if record.approval_number is not None:
        assert record.approval_number not in amount_values
    assert "APPROVAL-TEST-42" not in amount_values
    assert "APPROVAL-42" not in amount_values
    assert 1 not in amount_values
    assert 3 not in amount_values
    assert 10 not in amount_values
    assert 88 not in amount_values


def assert_downstream_receipt_shape(record: ReceiptRecord) -> None:
    data = record.to_dict()
    required_keys = {
        "id",
        "source_file_id",
        "source_file_name",
        "transaction_date",
        "merchant_name",
        "total_amount",
        "supply_amount",
        "vat_amount",
        "review_status",
        "parser_warnings",
        "validation_warnings",
        "amount_warnings",
        "vat_warnings",
        "sensitive_text_detected",
        "sensitive_text_categories",
        "card_match_status",
        "ocr_provider",
    }
    assert required_keys.issubset(data)
