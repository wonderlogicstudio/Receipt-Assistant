from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.card_models import CardTransaction
from app.core.enums import CardMatchStatus, ReviewStatus
from app.core.ocr_models import OCRResult
from app.core.receipt_models import ReceiptRecord


SHEET_SUMMARY = "01_정리내역"
SHEET_MONTHLY = "02_월별요약"
SHEET_CATEGORY = "03_비용항목별요약"
SHEET_REVIEW = "04_검토필요"
SHEET_CARD_MATCH = "05_카드대조"
SHEET_SOURCE_FILES = "06_원본파일목록"
SHEET_PROCESS_LOG = "07_처리로그"
SHEET_EXCLUDED = "08_제외내역"

AMOUNT_HEADERS = {
    "총금액",
    "할인전금액",
    "할인금액",
    "최종결제금액",
    "실제결제금액",
    "공급가액",
    "부가세",
    "추정공급가액",
    "추정부가세",
    "총금액합계",
    "공급가액합계",
    "부가세합계",
    "영수증금액",
    "카드금액",
}
DATE_HEADERS = {"사용일자", "영수증사용일자", "카드사용일자"}
VAT_REFERENCE_HEADERS = {"추정공급가액", "추정부가세", "부가세출처", "부가세검토사유"}
FORMULA_PREFIXES = ("=", "+", "-", "@")


@dataclass(frozen=True)
class ProcessingLogEntry:
    step: str
    status: str
    message: str
    related_file_name: str | None = None
    created_at: datetime | None = None


class ExcelExporter:
    sheet_names = [
        SHEET_SUMMARY,
        SHEET_MONTHLY,
        SHEET_CATEGORY,
        SHEET_REVIEW,
        SHEET_CARD_MATCH,
        SHEET_SOURCE_FILES,
        SHEET_PROCESS_LOG,
        SHEET_EXCLUDED,
    ]
    summary_sheet_name = SHEET_SUMMARY
    monthly_sheet_name = SHEET_MONTHLY
    review_sheet_name = SHEET_REVIEW
    raw_ocr_sheet_name = SHEET_PROCESS_LOG

    summary_headers = [
        "No",
        "사용일자",
        "상호",
        "사업자번호",
        "총금액",
        "할인전금액",
        "할인금액",
        "최종결제금액",
        "실제결제금액",
        "금액판단기준",
        "금액검토사유",
        "공급가액",
        "부가세",
        "추정공급가액",
        "추정부가세",
        "부가세출처",
        "부가세검토사유",
        "민감문구감지",
        "민감문구유형",
        "결제수단",
        "시스템추천비용항목",
        "최종비용항목",
        "검수상태",
        "검토사유",
        "카드매칭상태",
        "카드매칭사유",
        "중복의심여부",
        "원본파일명",
        "비고",
    ]

    monthly_headers = ["월", "건수", "총금액합계", "공급가액합계", "부가세합계", "검토필요건수", "카드미매칭건수"]
    category_headers = ["비용항목", "건수", "총금액합계", "공급가액합계", "부가세합계", "검토필요건수"]
    review_headers = [
        "No",
        "사용일자",
        "상호",
        "총금액",
        "검수상태",
        "검토사유",
        "경고",
        "카드매칭상태",
        "카드매칭사유",
        "중복의심여부",
        "원본파일명",
    ]
    card_match_headers = [
        "No",
        "영수증사용일자",
        "영수증상호",
        "영수증금액",
        "카드사용일자",
        "카드가맹점명",
        "카드금액",
        "카드명",
        "승인번호",
        "매칭상태",
        "매칭점수",
        "매칭사유",
        "후보건수",
        "원본파일명",
    ]
    source_file_headers = ["No", "원본파일명", "파일유형", "처리유형", "처리상태", "건수", "경고", "내부추적ID"]
    process_log_headers = ["No", "단계", "상태", "메시지", "관련파일명", "생성시각"]
    excluded_headers = ["No", "사용일자", "상호", "총금액", "제외사유", "원본파일명", "비고"]

    def export(
        self,
        receipts: Iterable[ReceiptRecord],
        card_transactions: Iterable[CardTransaction] | None = None,
        output_path: str | Path | None = None,
    ) -> bytes | str:
        workbook = self.build_workbook(receipts, card_transactions)
        if output_path is None:
            buffer = BytesIO()
            workbook.save(buffer)
            return buffer.getvalue()

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return str(path)

    def export_to_bytes(
        self,
        receipts: Iterable[ReceiptRecord],
        card_transactions: Iterable[CardTransaction] | Iterable[OCRResult] | None = None,
        ocr_results: Iterable[OCRResult] | None = None,
    ) -> bytes:
        cards, ocr_items = self._split_legacy_inputs(card_transactions, ocr_results)
        return self._workbook_to_bytes(self.build_workbook(receipts, cards, ocr_results=ocr_items))

    def build_workbook(
        self,
        receipts: Iterable[ReceiptRecord],
        card_transactions: Iterable[CardTransaction] | None = None,
        *,
        processing_logs: Iterable[ProcessingLogEntry] | None = None,
        ocr_results: Iterable[OCRResult] | None = None,
    ) -> Workbook:
        receipt_list = list(receipts)
        card_transaction_list = list(card_transactions or [])
        processing_log_list = list(processing_logs or [])
        ocr_result_list = list(ocr_results or [])

        workbook = Workbook()
        workbook.remove(workbook.active)

        included_receipts = [receipt for receipt in receipt_list if not self._is_excluded(receipt)]
        excluded_receipts = [receipt for receipt in receipt_list if self._is_excluded(receipt)]
        card_by_id = {transaction.id: transaction for transaction in card_transaction_list}

        self._add_summary_sheet(workbook, included_receipts)
        self._add_monthly_sheet(workbook, included_receipts)
        self._add_category_sheet(workbook, included_receipts)
        self._add_review_sheet(workbook, included_receipts)
        self._add_card_match_sheet(workbook, receipt_list, card_by_id)
        self._add_source_file_sheet(workbook, receipt_list, card_transaction_list)
        self._add_process_log_sheet(workbook, receipt_list, card_transaction_list, processing_log_list, ocr_result_list)
        self._add_excluded_sheet(workbook, excluded_receipts)
        return workbook

    def make_filename(self, now: datetime | None = None) -> str:
        timestamp = (now or datetime.now()).strftime("%Y%m%d_%H%M%S")
        return f"receipt_report_{timestamp}.xlsx"

    def _add_summary_sheet(self, workbook: Workbook, receipts: list[ReceiptRecord]) -> None:
        rows = [self._summary_row(index, receipt) for index, receipt in enumerate(receipts, start=1)]
        self._append_sheet(workbook, SHEET_SUMMARY, self.summary_headers, rows)

    def _add_monthly_sheet(self, workbook: Workbook, receipts: list[ReceiptRecord]) -> None:
        totals: dict[str, dict[str, int]] = defaultdict(
            lambda: {
                "count": 0,
                "total": 0,
                "supply": 0,
                "vat": 0,
                "needs_review": 0,
                "card_unmatched": 0,
            }
        )
        for receipt in receipts:
            month = receipt.transaction_date.strftime("%Y-%m") if receipt.transaction_date else "날짜미확인"
            totals[month]["count"] += 1
            totals[month]["total"] += receipt.total_amount or 0
            totals[month]["supply"] += receipt.supply_amount or 0
            totals[month]["vat"] += receipt.vat_amount or 0
            if self._is_review_needed(receipt):
                totals[month]["needs_review"] += 1
            if receipt.card_match_status in {CardMatchStatus.UNMATCHED.value, CardMatchStatus.CANDIDATE.value}:
                totals[month]["card_unmatched"] += 1

        rows = [
            [
                month,
                item["count"],
                item["total"],
                item["supply"],
                item["vat"],
                item["needs_review"],
                item["card_unmatched"],
            ]
            for month, item in sorted(totals.items())
        ]
        self._append_sheet(workbook, SHEET_MONTHLY, self.monthly_headers, rows)

    def _add_category_sheet(self, workbook: Workbook, receipts: list[ReceiptRecord]) -> None:
        totals: dict[str, dict[str, int]] = defaultdict(
            lambda: {"count": 0, "total": 0, "supply": 0, "vat": 0, "needs_review": 0}
        )
        for receipt in receipts:
            category = receipt.expense_category.strip() if receipt.expense_category else "미확정"
            totals[category]["count"] += 1
            totals[category]["total"] += receipt.total_amount or 0
            totals[category]["supply"] += receipt.supply_amount or 0
            totals[category]["vat"] += receipt.vat_amount or 0
            if self._is_review_needed(receipt):
                totals[category]["needs_review"] += 1

        rows = [
            [category, item["count"], item["total"], item["supply"], item["vat"], item["needs_review"]]
            for category, item in sorted(totals.items())
        ]
        self._append_sheet(workbook, SHEET_CATEGORY, self.category_headers, rows)

    def _add_review_sheet(self, workbook: Workbook, receipts: list[ReceiptRecord]) -> None:
        review_receipts = [receipt for receipt in receipts if self._is_review_needed(receipt)]
        rows = [
            [
                index,
                receipt.transaction_date,
                receipt.merchant_name,
                receipt.total_amount,
                receipt.review_status,
                self._review_reason_text(receipt),
                self._warning_text(receipt),
                receipt.card_match_status,
                receipt.card_match_reason,
                self._has_duplicate_signal(receipt),
                self._safe_file_name(receipt.source_file_name),
            ]
            for index, receipt in enumerate(review_receipts, start=1)
        ]
        self._append_sheet(workbook, SHEET_REVIEW, self.review_headers, rows)

    def _add_card_match_sheet(
        self,
        workbook: Workbook,
        receipts: list[ReceiptRecord],
        card_by_id: dict[str, CardTransaction],
    ) -> None:
        rows = []
        for index, receipt in enumerate(receipts, start=1):
            transaction = card_by_id.get(receipt.matched_card_transaction_id or "")
            rows.append(
                [
                    index,
                    receipt.transaction_date,
                    receipt.merchant_name,
                    receipt.total_amount,
                    transaction.transaction_date if transaction else None,
                    transaction.merchant_name if transaction else None,
                    transaction.total_amount if transaction else None,
                    transaction.card_name if transaction else None,
                    transaction.approval_number if transaction else None,
                    receipt.card_match_status,
                    receipt.card_match_score,
                    receipt.card_match_reason,
                    receipt.card_match_candidate_count,
                    self._safe_file_name(receipt.source_file_name),
                ]
            )
        self._append_sheet(workbook, SHEET_CARD_MATCH, self.card_match_headers, rows)

    def _add_source_file_sheet(
        self,
        workbook: Workbook,
        receipts: list[ReceiptRecord],
        card_transactions: list[CardTransaction],
    ) -> None:
        source_rows: dict[tuple[str, str], dict[str, object]] = {}
        for receipt in receipts:
            key = ("receipt", receipt.source_file_id)
            row = source_rows.setdefault(
                key,
                {
                    "name": self._safe_file_name(receipt.source_file_name),
                    "file_type": _extension(receipt.source_file_name),
                    "process_type": "영수증",
                    "status": "processed",
                    "count": 0,
                    "warnings": [],
                    "source_file_id": receipt.source_file_id,
                },
            )
            row["count"] = int(row["count"]) + 1
            warnings = row["warnings"]
            if isinstance(warnings, list):
                warnings.extend(receipt.parser_warnings)
                warnings.extend(receipt.validation_warnings)

        for transaction in card_transactions:
            key = ("card", transaction.source_file_id)
            row = source_rows.setdefault(
                key,
                {
                    "name": self._safe_file_name(transaction.source_file_name),
                    "file_type": _extension(transaction.source_file_name),
                    "process_type": "카드사용내역",
                    "status": "processed",
                    "count": 0,
                    "warnings": [],
                    "source_file_id": transaction.source_file_id,
                },
            )
            row["count"] = int(row["count"]) + 1
            warnings = row["warnings"]
            if isinstance(warnings, list):
                warnings.extend(transaction.warning_messages)

        rows = []
        for index, row in enumerate(source_rows.values(), start=1):
            rows.append(
                [
                    index,
                    row["name"],
                    row["file_type"],
                    row["process_type"],
                    row["status"],
                    row["count"],
                    self._join_unique(row["warnings"] if isinstance(row["warnings"], list) else []),
                    row["source_file_id"],
                ]
            )
        self._append_sheet(workbook, SHEET_SOURCE_FILES, self.source_file_headers, rows)

    def _add_process_log_sheet(
        self,
        workbook: Workbook,
        receipts: list[ReceiptRecord],
        card_transactions: list[CardTransaction],
        processing_logs: list[ProcessingLogEntry],
        ocr_results: list[OCRResult],
    ) -> None:
        rows = []
        for index, log in enumerate(processing_logs, start=1):
            rows.append(
                [
                    index,
                    log.step,
                    log.status,
                    log.message,
                    self._safe_file_name(log.related_file_name),
                    log.created_at or datetime.now(),
                ]
            )

        next_index = len(rows) + 1
        if not rows:
            for receipt in receipts:
                rows.append(
                    [
                        next_index,
                        "receipt_export",
                        receipt.review_status,
                        self._short_message(receipt),
                        self._safe_file_name(receipt.source_file_name),
                        receipt.created_at,
                    ]
                )
                next_index += 1
            for transaction in card_transactions:
                rows.append(
                    [
                        next_index,
                        "card_statement_parse",
                        transaction.parse_status,
                        self._join_unique(transaction.warning_messages) or "processed",
                        self._safe_file_name(transaction.source_file_name),
                        transaction.created_at,
                    ]
                )
                next_index += 1
            for result in ocr_results:
                rows.append(
                    [
                        next_index,
                        "ocr_extract",
                        result.status,
                        result.error_message or "processed",
                        self._safe_file_name(result.source_file_name),
                        result.created_at,
                    ]
                )
                next_index += 1

        self._append_sheet(workbook, SHEET_PROCESS_LOG, self.process_log_headers, rows)

    def _add_excluded_sheet(self, workbook: Workbook, receipts: list[ReceiptRecord]) -> None:
        rows = [
            [
                index,
                receipt.transaction_date,
                receipt.merchant_name,
                receipt.total_amount,
                self._review_reason_text(receipt) or "제외",
                self._safe_file_name(receipt.source_file_name),
                receipt.user_note,
            ]
            for index, receipt in enumerate(receipts, start=1)
        ]
        self._append_sheet(workbook, SHEET_EXCLUDED, self.excluded_headers, rows)

    def _summary_row(self, index: int, receipt: ReceiptRecord) -> list[object]:
        return [
            index,
            receipt.transaction_date,
            receipt.merchant_name,
            receipt.business_number,
            receipt.total_amount,
            receipt.gross_amount,
            receipt.discount_amount,
            receipt.final_amount,
            receipt.paid_amount,
            receipt.amount_source,
            self._join_unique(receipt.amount_warnings),
            receipt.supply_amount,
            receipt.vat_amount,
            receipt.estimated_supply_amount,
            receipt.estimated_vat_amount,
            receipt.vat_source,
            self._join_unique([*receipt.vat_warnings, receipt.vat_note] if receipt.vat_note else receipt.vat_warnings),
            "Y" if receipt.sensitive_text_detected else "",
            self._join_unique(receipt.sensitive_text_categories),
            receipt.payment_method,
            receipt.recommended_expense_category,
            receipt.expense_category,
            receipt.review_status,
            self._review_reason_text(receipt),
            receipt.card_match_status,
            receipt.card_match_reason,
            self._has_duplicate_signal(receipt),
            self._safe_file_name(receipt.source_file_name),
            receipt.user_note,
        ]

    def _append_sheet(self, workbook: Workbook, title: str, headers: list[str], rows: Sequence[Sequence[object]]) -> None:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append([self._cell_value(value) for value in row])
        self._format_sheet(sheet)

    def _format_sheet(self, sheet) -> None:
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        reference_fill = PatternFill("solid", fgColor="FFF2CC")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = reference_fill if cell.value in VAT_REFERENCE_HEADERS else header_fill
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            column_letter = get_column_letter(column_cells[0].column)
            sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                header = sheet.cell(row=1, column=cell.column).value
                if header in AMOUNT_HEADERS:
                    cell.number_format = "#,##0"
                if header in DATE_HEADERS:
                    cell.number_format = "yyyy-mm-dd"
                if header in VAT_REFERENCE_HEADERS:
                    cell.fill = reference_fill

    def _cell_value(self, value: object) -> object:
        if isinstance(value, Path):
            return self._safe_file_name(str(value))
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            return self._safe_text(value)
        return value

    def _safe_text(self, value: str) -> str:
        text = value.strip()
        text = self._safe_file_name(text) if _looks_like_path(text) else text
        if text.startswith(FORMULA_PREFIXES):
            return f"'{text}"
        return text

    def _safe_file_name(self, value: str | None) -> str | None:
        if not value:
            return None
        name = Path(str(value).replace("\\", "/")).name
        if name.startswith(FORMULA_PREFIXES):
            return f"'{name}"
        return name

    def _review_reason_text(self, receipt: ReceiptRecord) -> str:
        values = [
            *self._split_text(receipt.review_reason),
            *receipt.review_reasons,
            *receipt.parser_warnings,
            *receipt.validation_warnings,
            *receipt.amount_warnings,
            *receipt.vat_warnings,
            *receipt.sensitive_warnings,
        ]
        return self._join_unique(values)

    def _warning_text(self, receipt: ReceiptRecord) -> str:
        return self._join_unique(
            [
                *receipt.parser_warnings,
                *receipt.validation_warnings,
                *receipt.amount_warnings,
                *receipt.vat_warnings,
                *receipt.sensitive_warnings,
            ]
        )

    def _short_message(self, receipt: ReceiptRecord) -> str:
        reasons = self._review_reason_text(receipt)
        if reasons:
            return reasons
        if self._is_excluded(receipt):
            return "excluded"
        return "processed"

    def _is_review_needed(self, receipt: ReceiptRecord) -> bool:
        if receipt.review_status == ReviewStatus.NEEDS_REVIEW.value:
            return True
        if receipt.review_reasons or receipt.validation_warnings or receipt.parser_warnings:
            return True
        if any(warning != "discount_detected" for warning in receipt.amount_warnings):
            return True
        if receipt.vat_review_required or receipt.vat_warnings:
            return True
        if receipt.sensitive_text_detected or receipt.sensitive_warnings:
            return True
        if receipt.card_match_status in {CardMatchStatus.UNMATCHED.value, CardMatchStatus.CANDIDATE.value}:
            return True
        return self._has_duplicate_signal(receipt)

    def _is_excluded(self, receipt: ReceiptRecord) -> bool:
        return receipt.review_status == ReviewStatus.EXCLUDED.value or getattr(receipt, "export_included", True) is False

    def _has_duplicate_signal(self, receipt: ReceiptRecord) -> bool:
        return bool(receipt.duplicate_group_id or receipt.duplicate_candidate_ids or receipt.duplicate_reason)

    def _join_unique(self, values: list[str]) -> str:
        return "; ".join(_unique([str(value).strip() for value in values if value]))

    def _split_text(self, value: str | None) -> list[str]:
        if not value:
            return []
        return [item.strip() for item in value.split(";") if item.strip()]

    def _split_legacy_inputs(
        self,
        card_transactions: Iterable[CardTransaction] | Iterable[OCRResult] | None,
        ocr_results: Iterable[OCRResult] | None,
    ) -> tuple[list[CardTransaction], list[OCRResult]]:
        items = list(card_transactions or [])
        if items and isinstance(items[0], OCRResult):
            return [], list(items) + list(ocr_results or [])
        return list(items), list(ocr_results or [])

    def _workbook_to_bytes(self, workbook: Workbook) -> bytes:
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()


def _extension(file_name: str | None) -> str:
    if not file_name:
        return ""
    return Path(str(file_name).replace("\\", "/")).suffix.lower().lstrip(".")


def _looks_like_path(value: str) -> bool:
    return ":\\" in value or ":/" in value or "\\" in value or "/" in value


def _unique(values: list[str]) -> list[str]:
    result: list[str] = []
    for value in values:
        if value and value not in result:
            result.append(value)
    return result
